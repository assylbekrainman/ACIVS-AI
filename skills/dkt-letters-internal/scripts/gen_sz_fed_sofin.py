"""
gen_sz_fed_sofin.py — Служебная записка в ФЭД АО «Фонд науки»
О проверке подтверждающих документов по исполнению обязательств по со-финансированию

ПАРАМЕТРЫ ПРОЕКТА — заполнить перед запуском:
"""

# ─────────────────── ПАРАМЕТРЫ ПРОЕКТА ────────────────────────────────────
IRN             = "DP23692342"
PROJECT_NAME    = (
    "«Коммерциализация инновационных технологий производства "
    "функциональных овощных соков с высоким содержанием полезных компонентов»"
)
GP_NAME_SHORT   = "СПК «MELON Juice Co»"
GP_NAME_FULL    = (
    "Сельскохозяйственным производственным кооперативом «MELON Juice Co»"
)
CONTRACT_NO     = "№107"
CONTRACT_DATE   = "24.09.2024 г."
STAGE           = "2"
GP_LETTER_NO    = "№12.05-26-1"
GP_LETTER_DATE  = "12.05.2026 г."
SOFIN_AMOUNT    = "39 490 300,00"
SOFIN_AMOUNT_WORDS = (
    "тридцать девять миллионов четыреста девяносто тысяч триста"
)
FED_DIRECTOR    = "_________________________"
DKT_SIGNER      = "Директор ДКТ"
DKT_SIGNER_NAME = "_____________________"
EXECUTOR        = "______________________"
EXECUTOR_PHONE  = "_____________"
OUTPUT_PATH     = f"СЗ_ФЭД_Софинансирование_{IRN}.docx"
# ──────────────────────────────────────────────────────────────────────────

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
section = doc.sections[0]
section.top_margin    = Cm(2.0)
section.bottom_margin = Cm(2.0)
section.left_margin   = Cm(3.0)
section.right_margin  = Cm(1.5)

def _set_font(run, size=12):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    r = run._r
    rPr = r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rFonts.set(qn(attr), 'Times New Roman')
    existing = rPr.find(qn('w:rFonts'))
    if existing is not None:
        rPr.remove(existing)
    rPr.insert(0, rFonts)

def add_para(doc, text, bold=False, italic=False, size=12,
             align=WD_ALIGN_PARAGRAPH.LEFT,
             space_before=0, space_after=6, first_line_indent=None):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after  = Pt(space_after)
    if first_line_indent is not None:
        pf.first_line_indent = Cm(first_line_indent)
    run = p.add_run(text)
    run.bold = bold; run.italic = italic
    _set_font(run, size)
    return p

def add_mixed(doc, parts, size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
              space_before=0, space_after=6, first_line_indent=1.25):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after  = Pt(space_after)
    if first_line_indent is not None:
        pf.first_line_indent = Cm(first_line_indent)
    for text, bold, italic in parts:
        run = p.add_run(text)
        run.bold = bold; run.italic = italic
        _set_font(run, size)
    return p

# АДРЕСАТ
p_addr = doc.add_paragraph()
p_addr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p_addr.paragraph_format.space_after = Pt(0)
for line in ["Директору Финансово-экономического", "департамента АО «Фонд науки»"]:
    r = p_addr.add_run(line + "\n"); _set_font(r)
p_addr2 = doc.add_paragraph()
p_addr2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p_addr2.paragraph_format.space_after = Pt(18)
r2 = p_addr2.add_run(FED_DIRECTOR); _set_font(r2)

# ЗАГОЛОВОК
add_para(doc, "СЛУЖЕБНАЯ ЗАПИСКА", bold=True, size=13,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=4)
add_para(doc,
         "О проверке подтверждающих документов по исполнению\n"
         "обязательств по со-финансированию",
         size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=10)
add_para(doc, f"по проекту ИРН {IRN}", bold=True, size=12,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=20)

# ТЕКСТ
add_mixed(doc, [
    ("Грантополучателем по проекту ИРН ", False, False),
    (IRN, True, False),
    (f" {PROJECT_NAME} — ", False, False),
    (GP_NAME_FULL, True, False),
    (f" — в соответствии с Договором {CONTRACT_NO} от {CONTRACT_DATE} "
     f"предоставлен пакет документов по исполнению обязательств "
     f"по со-финансированию за {STAGE}-й этап реализации Проекта.", False, False),
], space_after=10)

add_mixed(doc, [
    (f"Письмом исх. {GP_LETTER_NO} от {GP_LETTER_DATE} Грантополучатель "
     f"подтвердил, что обязательства по освоению средств со-финансирования "
     f"исполнены в полном объёме на сумму ", False, False),
    (f"{SOFIN_AMOUNT} ({SOFIN_AMOUNT_WORDS}) тенге", True, False),
    (". К письму приложены подтверждающие документы и выписка о движении "
     "средств со-финансирования с отметкой банка за отчётный период.", False, False),
], space_after=10)

add_para(doc,
    f"В связи с вышеизложенным, просим Вас рассмотреть и проверить "
    f"предоставленные подтверждающие документы по исполнению обязательств "
    f"по со-финансированию Грантополучателя {GP_NAME_SHORT} "
    f"по {STAGE}-му этапу Проекта ИРН {IRN} и направить заключение "
    f"в Департамент коммерциализации технологий.",
    size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    first_line_indent=1.25, space_before=0, space_after=20)

# ПРИЛОЖЕНИЯ
add_para(doc, "Приложения:", bold=True, size=12,
         align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4)
annexes = [
    f"Письмо {GP_NAME_SHORT} исх. {GP_LETTER_NO} от {GP_LETTER_DATE};",
    f"Подтверждающие документы по исполнению обязательств по со-финансированию "
    f"на сумму {SOFIN_AMOUNT} тенге;",
    "Выписка о движении средств со-финансирования с отметкой банка "
    "за отчётный период по счёту со-финансирования.",
]
for txt in annexes:
    p = doc.add_paragraph(style='List Number')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(1.25)
    run = p.add_run(txt); _set_font(run)

# ПОДПИСЬ
doc.add_paragraph()
tbl = doc.add_table(rows=2, cols=2)
tbl.style = 'Table Grid'
tblPr = tbl._tbl.find(qn('w:tblPr'))
if tblPr is None:
    tblPr = OxmlElement('w:tblPr'); tbl._tbl.insert(0, tblPr)
tblBorders = OxmlElement('w:tblBorders')
for edge in ('top','left','bottom','right','insideH','insideV'):
    el = OxmlElement(f'w:{edge}'); el.set(qn('w:val'), 'none'); tblBorders.append(el)
tblPr.append(tblBorders)

def cell_text(cell, text, bold=False, size=12, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.paragraphs[0].clear()
    p = cell.paragraphs[0]; p.alignment = align
    run = p.add_run(text); run.bold = bold; _set_font(run, size)

cell_text(tbl.rows[0].cells[0], DKT_SIGNER)
cell_text(tbl.rows[0].cells[1], f"_________________ / {DKT_SIGNER_NAME}")
cell_text(tbl.rows[1].cells[0], "")
cell_text(tbl.rows[1].cells[1], "(подпись)              (ФИО)", size=10)

doc.add_paragraph()
add_para(doc, f"Исп.: {EXECUTOR} / Тел.: {EXECUTOR_PHONE}",
         size=10, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=0)

doc.save(OUTPUT_PATH)
print(f"Сохранено: {OUTPUT_PATH}")
