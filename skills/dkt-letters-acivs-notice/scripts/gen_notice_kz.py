"""
gen_notice_kz.py — Генератор хабарламы ГА БҚМПТ бойынша (Тип 10-KZ)
АО «Фонд науки» / ДКТ

Бұрын бұл скрипт `dkt-letters` скиллінде `references/`-те аталған, бірақ
нақты файл болмаған (SKILL.md оны сілтеп тұрды, ал скрипт өзі жоқ болатын —
Фаза 4 аудитінде табылған нақты олқылық). Бұл файл сол олқылықты жабады —
Тип 10 (орысша, `gen_acivs_notice.py`) құрылымының дәл сол логикасын
қазақша нұсқасы үшін қайталайды, тек мәтін мен сценарий тармақтары
(ШАГ 3, `SKILL.md`-дегі Сценарий А–Д) қазақша.

**МАҢЫЗДЫ:** параграф индекстері (P[4], P[8]...) `gen_acivs_notice.py`-дегі
сияқты орысша үлгінің құрылымын болжайды («уведомлениенің алдыңғы
нұсқасы» — SKILL.md-де айтылғандай, логотип+адрес блогы сақталған қазақша
аудармасы негізгі үлгі ретінде қолданылады). **Бірінші рет пайдаланар
алдында нақты үлгі файлымен индекстерді тексеріп ал** — олар нақты
құжаттың абзац құрылымына байланысты өзгеруі мүмкін, дәл осы ескерту
`gen_acivs_notice.py`-дің өзінде де бар.

Тәуелділік: pip install python-docx --break-system-packages
"""

import copy, glob
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt
from docx.text.paragraph import Paragraph

# ===================================================================
#  ЖОБА ПАРАМЕТРЛЕРІ — Шарттан/БҚМПТ қорытындысынан толтыру
# ===================================================================

IRN           = "DP23692342"
STAGE         = 2
STAGE_KZ      = "2-ші"
PERIOD        = "2025 жылғы қаңтар-желтоқсан"

# Есеп түрі: True — аралық (келесі кезең бар), False — қорытынды (соңғы кезең)
IS_INTERIM = True

# Алушы (оң жақ блок, қалың)
GA_LINE1 = "«MELON juice Co» ЖШС"
GA_LINE2 = ""   # "" егер бір жолда болса

# Шарт деректері
CONTRACT_NUM  = "107"
CONTRACT_DATE = "«24» қыркүйек 2024 жыл"
PROJECT_NAME  = (
    "«Пайдалы құрамдас бөліктердің жоғары мөлшерімен функционалды "
    "көкөніс шырындарын өндірудің инновациялық технологияларын "
    "коммерцияландыру»"
)

# Шарт тармақтары (нақты Шарттан алу керек!)
CLAUSE_NEXT_STAGE = "3.4"   # келесі кезеңді қаржыландыру құқығы (тек аралық үшін)
CLAUSE_ECONOMY    = "3.7"   # үнемдеуді келесі кезең қаржыландыруынан ұстау

# Қолтаңбашы
SIGNATORY_TITLE = "Басқарушы директор"
SIGNATORY_NAME  = "А. Рахимжанов"

# Орындаушы
EXECUTOR_NAME   = "Орынд.: А. Снадин"
EXECUTOR_PHONE  = "8 (7172) 57-50-08 / 57-50-15"

# ===================================================================
#  СЦЕНАРИЙ ТАҢДАУ (SKILL.md ШАГ 3-тегі А–Д сценарийлеріне сәйкес)
# ===================================================================

# --- Б: үнемдеу (None немесе 0 — үнемдеу жоқ) ---
ECO_SUM_KZ = (
    "3 158 414 (үш миллион бір жүз елу сегіз мың төрт жүз "
    "он төрт) теңге 00 тиын"
)  # None немесе "" — үнемдеу жоқ дегенді білдіреді

# --- В: бірлескен қаржыландыру жетіспеушілігі ---
SOFIN_SHORTFALL = True
SOFIN_ACCOUNT = "KZ558562203140010240"
SOFIN_BANK    = "АО «БанкЦентрКредит»"
SOFIN_NEDOSTACH_KZ = (
    "39 490 300,00 (отыз тоғыз миллион төрт жүз тоқсан мың үш жүз) теңге"
)
CLAUSE_SOFIN_DEPOSIT  = "2.16"   # р/с-қа аудару міндеті
CLAUSE_SOFIN_SPENDING = "2.8"    # тек Жоба мақсатына жұмсау

# --- Г: КП іс-шаралары орындалмаған (None — барлығы орындалды) ---
KP_NOT_DONE_ITEM = None  # мысалы: "3.2-тармағында көзделген жабдықты сатып алу"
CLAUSE_KP = "3.2"

# --- Д: мақсатсыз/жоспарланбаған пайдалану (None — анықталмады) ---
MISUSE_SUM_KZ = None  # мысалы: "1 200 000 (бір миллион екі жүз мың) теңге"
CLAUSE_MISUSE = "3.9"

# Счета ГА (В сценарийінде үзінді сұрау үшін)
ACCOUNTS = [
    {"type": "эскроу-шот", "num": "KZ168562203140864348", "bank": "АО «БанкЦентрКредит»"},
    {"type": "ағымдағы шот", "num": "KZ668562203141697551", "bank": "АО «БанкЦентрКредит»"},
    {"type": "бірлескен қаржыландыру шоты", "num": SOFIN_ACCOUNT, "bank": SOFIN_BANK},
]

# БҚМПТ
ACIVS_PAGES = 13

# Жолдар
TEMPLATE = ""
OUTPUT   = f"Хабарлама_БҚМПТ{STAGE}кезең_{IRN}.docx"

if not TEMPLATE:
    candidates = (
        glob.glob("/sessions/*/mnt/uploads/Хабарлама_БҚМПТ*кезең_*.docx") +
        glob.glob("/sessions/*/mnt/**/*Хабарлама_БҚМПТ*кезең*.docx", recursive=True) +
        # Резервный вариант: РУ-шаблон уведомления (та же структура логотип+адрес)
        glob.glob("/sessions/*/mnt/uploads/Уведомление_АЦИВС*этап_*.docx") +
        glob.glob("/sessions/*/mnt/**/*Уведомление_АЦИВС*этап*.docx", recursive=True)
    )
    if not candidates:
        raise FileNotFoundError(
            "Үлгі табылмады. Хабарлама_БҚМПТ*кезең_*.docx жүктеңіз, немесе "
            "SKILL.md-де айтылғандай — Тип 10 (орысша) шығысын үлгі ретінде "
            "қолданыңыз (логотип+адрес блогы бірдей)."
        )
    TEMPLATE = candidates[0]
    print(f"Үлгі: {TEMPLATE}")

# ===================================================================
#  КӨМЕКШІ ФУНКЦИЯЛАР (gen_acivs_notice.py-мен бірдей)
# ===================================================================

def clear_runs(para):
    p = para._p
    for tag in ("w:r", "w:ins", "w:hyperlink", "w:del"):
        for el in p.findall(qn(tag)):
            p.remove(el)

def add_run(para, text, bold=False, italic=False, sz=12):
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(sz)
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rFonts.set(qn(attr), "Times New Roman")
    return run

def clone_ppr(para):
    pPr = para._p.find(qn("w:pPr"))
    return copy.deepcopy(pPr) if pPr is not None else None

def insert_after(ref_para, text_parts, pPr_clone=None):
    new_p = OxmlElement("w:p")
    if pPr_clone is not None:
        new_p.append(copy.deepcopy(pPr_clone))
    ref_para._p.addnext(new_p)
    para = Paragraph(new_p, ref_para._p.getparent())
    for part in text_parts:
        add_run(para, part["text"],
                bold=part.get("bold", False),
                italic=part.get("italic", False),
                sz=part.get("sz", 12))
    return para

# ===================================================================
#  МӘТІН БЛОКТАРЫ (SKILL.md ШАГ 2-3 бойынша)
# ===================================================================

report_kind = "аралық" if IS_INTERIM else "қорытынды (соңғы)"

PARA_INTRO = (
    f"«Ғылым қоры» АҚ (бұдан әрі – Қор) №{IRN} {PROJECT_NAME} жобасы "
    f"(бұдан әрі – Жоба) бойынша {STAGE_KZ} кезеңнің {report_kind} есебіне "
    f"бөлінген қаражаттың мақсатты пайдаланылуын талдау қорытындысын "
    f"(бұдан әрі – БҚМПТ) таныстыру және қол қою үшін жолдайды. Жоба "
    f"{GA_LINE1}-мен Ғылыми және (немесе) ғылыми-техникалық қызмет "
    f"нәтижелерін коммерцияландыруға грант беру туралы {CONTRACT_DATE} "
    f"№{CONTRACT_NUM} Шарт (бұдан әрі – Шарт) негізінде жүзеге асырылады."
)

if IS_INTERIM:
    PARA_FINANCING = (
        f"Шарттың {CLAUSE_NEXT_STAGE}-тармағына сәйкес, БҚМПТ оң қорытындысы "
        f"негізінде Грант алушы жобаның {STAGE + 1}-ші кезеңін қаржыландыру "
        f"құқығын алады."
    )
else:
    PARA_FINANCING = (
        f"{STAGE_KZ} кезең Жобаны іске асырудың қорытынды кезеңі болып "
        f"табылатындықтан, Шарт бойынша кейінгі кезеңді қаржыландыру "
        f"көзделмеген."
    )

has_economy = bool(ECO_SUM_KZ)
if not has_economy:
    PARA_ECONOMY = (
        f"Шарттың {CLAUSE_ECONOMY}-тармағына сәйкес, {STAGE_KZ} кезең "
        f"бойынша грант қаражатының үнемделуі 0 теңгені құрады; «Ғылым "
        f"қоры» АҚ есеп айырысу шотына қаражатты қайтару талап етілмейді."
    )
elif IS_INTERIM:
    PARA_ECONOMY = (
        f"Шарттың {CLAUSE_ECONOMY}-тармағына сәйкес, {STAGE_KZ} кезең "
        f"бойынша грант қаражатының үнемделуі {ECO_SUM_KZ} құрады; аталған "
        f"сома жобаның {STAGE + 1}-ші кезеңі қаржыландыру сомасынан "
        f"ұсталатын болады."
    )
else:
    PARA_ECONOMY = (
        f"Шарттың {CLAUSE_ECONOMY}-тармағына сәйкес, {STAGE_KZ} кезең "
        f"бойынша грант қаражатының үнемделуі {ECO_SUM_KZ} құрады; аталған "
        f"соманы «Ғылым қоры» АҚ есеп айырысу шотына қайтаруды сұраймыз."
    )

PARA_RESULTS_INTRO = "Сонымен, Жоба бойынша БҚМПТ қорытындысына сәйкес:"

# Bullet 1 — КП
if KP_NOT_DONE_ITEM:
    b1 = f"{STAGE_KZ} кезеңнің Күнтізбелік жоспарының {KP_NOT_DONE_ITEM} орындалмады;"
else:
    b1 = f"{STAGE_KZ} кезеңнің Күнтізбелік жоспарының барлық іс-шаралары толық көлемде орындалды;"

# Bullet 2 — үнемдеу
b2 = "грант қаражатының үнемделуі жоқ;" if not has_economy else f"грант қаражатының үнемделуі {ECO_SUM_KZ} құрады;"

# Bullet 3 — мақсатсыз пайдалану
if MISUSE_SUM_KZ:
    b3 = f"грант қаражатын мақсатсыз және жоспарланбаған пайдалану {MISUSE_SUM_KZ} сомасында анықталды;"
else:
    b3 = "грант қаражатын мақсатсыз және жоспарланбаған пайдалану анықталмады;"

# Bullet 4 — бірлескен қаржыландыру (соңғы бюллет — нүкте/нүктелі үтір SOFIN_SHORTFALL-ға байланысты)
b4_end = ";" if SOFIN_SHORTFALL else "."
if SOFIN_SHORTFALL:
    b4 = f"{STAGE_KZ} кезең бойынша бірлескен қаржыландыру міндеттемелері орындалмады{b4_end}"
else:
    b4 = f"{STAGE_KZ} кезең бойынша бірлескен қаржыландыру міндеттемелері толық орындалды{b4_end}"

REQUEST_INTRO = (
    "Жоғарыда баяндалғандарды негізге ала отырып, Қор осы хабарламаны "
    "жолдаған күннен бастап 5 (бес) жұмыс күні ішінде мынадай құжаттарды "
    "ұсынуды сұрайды:"
)
REQ_ITEM1 = f"{STAGE_KZ} кезеңнің {report_kind} есебі бойынша қол қойылған БҚМПТ қорытындысы."

PRILOZHENIE = (
    f"Қосымша: {IRN} жобасы бойынша {STAGE_KZ} кезеңнің {report_kind} "
    f"есебіне бөлінген қаражаттың мақсатты пайдалануын талдау қорытындысы "
    f"(БҚМПТ) {ACIVS_PAGES} парақта."
)

SIGNATORY_LINE = SIGNATORY_TITLE + " " * max(1, 60 - len(SIGNATORY_TITLE)) + SIGNATORY_NAME

# ===================================================================
#  ГЕНЕРАЦИЯ
# ===================================================================

doc   = Document(TEMPLATE)
paras = doc.paragraphs

# P[0] — логотип (тимеймiз); P[1] — мекенжай (тимеймiз)

clear_runs(paras[4]); add_run(paras[4], GA_LINE1, bold=True, sz=12)
if GA_LINE2:
    clear_runs(paras[5]); add_run(paras[5], GA_LINE2, bold=True, sz=12)

clear_runs(paras[8]);  add_run(paras[8],  PARA_INTRO,     sz=12)
clear_runs(paras[9]);  add_run(paras[9],  PARA_FINANCING, sz=12)
clear_runs(paras[10]); add_run(paras[10], PARA_ECONOMY,   sz=12)

pPr_body = clone_ppr(paras[8])

# P[11] — "Сонымен, ..." кіріспе, содан кейін P[12]-P[15] буллеттер деп
# болжанады (RU үлгісіндегі сияқты). Нақты индекстерді өз үлгіңізбен
# тексеріңіз — бұл RU нұсқасынан айырмашылық, себебі мұнда қосымша
# кіріспе абзацы (P[11]) бар.
if len(paras) > 11:
    clear_runs(paras[11]); add_run(paras[11], PARA_RESULTS_INTRO, sz=12)

bullets = [b1, b2, b3, b4]
bullet_start_idx = 12
for i, b in enumerate(bullets):
    idx = bullet_start_idx + i
    if idx < len(paras):
        clear_runs(paras[idx]); add_run(paras[idx], b, sz=12)
    else:
        # үлгіде жеткілікті буллет-абзац жоқ болса — соңғысынан кейін қосамыз
        ref = paras[-1]
        insert_after(ref, [{"text": b, "sz": 12}], pPr_clone=pPr_body)

all_paras = doc.paragraphs
p_req_idx = next((i for i, p in enumerate(all_paras)
                   if p.text.startswith("Жоғарыда баяндалғандарды")), None)
if p_req_idx is not None:
    clear_runs(all_paras[p_req_idx]); add_run(all_paras[p_req_idx], REQUEST_INTRO, sz=12)
    p_item1 = all_paras[p_req_idx + 1] if p_req_idx + 1 < len(all_paras) else None
    if p_item1 is not None:
        clear_runs(p_item1); add_run(p_item1, f"1) {REQ_ITEM1}", sz=12)
        last_item = p_item1
        if SOFIN_SHORTFALL:
            item2 = insert_after(last_item,
                [{"text": f"2) бірлескен қаржыландыру бойынша міндеттемелердің орындалғанын растайтын құжаттар;", "sz": 12}],
                pPr_clone=clone_ppr(last_item))
            accs = "; ".join(f"{a['type']} №{a['num']} {a['bank']}" for a in ACCOUNTS)
            insert_after(item2,
                [{"text": f"3) барлық шоттар бойынша қозғалыс туралы үзінді көшірмелер: {accs}.", "sz": 12}],
                pPr_clone=clone_ppr(last_item))

# Бірлескен қаржыландыру жетіспеушілігі бойынша 3 қосымша абзац (В сценарийі)
if SOFIN_SHORTFALL:
    all_paras = doc.paragraphs
    # соңғы буллеттен кейін, сұраныс абзацынан бұрын кірістіру
    anchor_idx = bullet_start_idx + len(bullets) - 1
    anchor = all_paras[anchor_idx] if anchor_idx < len(all_paras) else all_paras[-1]
    p1 = insert_after(anchor, [{"text":
        f"Шарттың {CLAUSE_SOFIN_DEPOSIT}-тармағына сәйкес, Грант алушы "
        f"бірлескен қаржыландыру қаражатын {SOFIN_BANK} {SOFIN_ACCOUNT} "
        f"есептік шотына жобада көзделген көлемде аударуға міндетті.", "sz": 12}],
        pPr_clone=pPr_body)
    p2 = insert_after(p1, [{"text":
        f"Шарттың {CLAUSE_SOFIN_SPENDING}-тармағына сәйкес, бірлескен "
        f"қаржыландыру қаражаты толық көлемде және тек Жобаның мақсаттары "
        f"мен міндеттеріне жұмсалуы тиіс.", "sz": 12}],
        pPr_clone=pPr_body)
    insert_after(p2, [{"text":
        f"БҚМПТ қорытындысына сәйкес {STAGE_KZ} кезең бойынша бірлескен "
        f"қаржыландыру міндеттемесінің орындалмауы {SOFIN_NEDOSTACH_KZ} "
        f"сомасында анықталды.", "sz": 12}],
        pPr_clone=pPr_body)

all_paras = doc.paragraphs
p_pril = next((p for p in all_paras if p.text.startswith("Қосымша:")), None)
if p_pril is not None:
    clear_runs(p_pril); add_run(p_pril, PRILOZHENIE, italic=True, sz=10)

all_paras = doc.paragraphs
p_sign = next((p for p in all_paras if "Басқарушы директор" in p.text or "Басқарма Төрағасы" in p.text), None)
if p_sign is not None:
    clear_runs(p_sign); add_run(p_sign, SIGNATORY_LINE, sz=12)

all_paras = doc.paragraphs
p_exec1 = next((p for p in all_paras if p.text.startswith("Орынд.:")), None)
if p_exec1 is not None:
    clear_runs(p_exec1); add_run(p_exec1, EXECUTOR_NAME, sz=10)
    idx = all_paras.index(p_exec1)
    if idx + 1 < len(all_paras):
        clear_runs(all_paras[idx + 1]); add_run(all_paras[idx + 1], EXECUTOR_PHONE, sz=10)

doc.save(OUTPUT)
print(f"Сақталды: {OUTPUT}")
print(
    "ЕСКЕРТУ: параграф индекстері (P[4],P[8]-P[15]...) шаблон құрылымына "
    "негізделген болжам — сақтағаннан кейін нәтижені PDF-ке айналдырып "
    "(soffie.py) көзбен тексеріңіз, әсіресе бірінші рет пайдаланғанда."
)
