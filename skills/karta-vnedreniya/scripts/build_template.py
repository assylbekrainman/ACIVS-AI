# -*- coding: utf-8 -*-
"""Одноразовый скрипт: из example_filled.docx делает чистый template_karta.docx."""
import os
from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, '..', 'references', 'example_filled.docx')
OUT = os.path.join(HERE, '..', 'references', 'template_karta.docx')

VAR_COLS = [3, 4, 5, 6, 7]  # План, Факт, Ответственный, Сроки, Примечание

def blank_cell_keep_format(cell):
    for p in cell.paragraphs:
        if p.runs:
            p.runs[0].text = ''
            for r in p.runs[1:]:
                r._element.getparent().remove(r._element)

d = Document(SRC)
t = d.tables[0]
cleared = []
for ri, row in enumerate(t.rows):
    if ri < 2:
        continue
    for ci in VAR_COLS:
        cell = row.cells[ci]
        if cell._tc in cleared:   # стабильный элемент lxml, не id()
            continue
        cleared.append(cell._tc)
        blank_cell_keep_format(cell)

# Конфигурируемые подпункты — название (col1) и форма завершения (col2) проектные,
# поэтому в каркасе их оставляем пустыми (заполняются скриптом по проекту).
CONFIG_NAME_NUMS = {'4.3.1', '4.3.2', '4.3.3', '6.1'}
for ri, row in enumerate(t.rows):
    num = row.cells[0].text.strip().rstrip('.')
    if num in CONFIG_NAME_NUMS:
        for ci in (1, 2):
            blank_cell_keep_format(row.cells[ci])

def set_para_text(p, text):
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        p.add_run(text)

for p in d.paragraphs:
    s = p.text.strip()
    if s.startswith('Грантополучатель'):
        set_para_text(p, 'Грантополучатель: {{GP}}')
    elif s.startswith('ИРН и наименование'):
        set_para_text(p, 'ИРН и наименование проекта: {{IRN_NAME}}')
    elif s.startswith('Договор'):
        set_para_text(p, 'Договор {{DOGOVOR}}')
    elif 'Пример (только для' in s:
        set_para_text(p, '')

d.save(OUT)
print('Saved template:', OUT)
