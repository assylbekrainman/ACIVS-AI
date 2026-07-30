# -*- coding: utf-8 -*-
"""
Заполнение «Карты внедрения по проекту коммерциализации РННТД» (Приложение №1),
форма АО «Фонд науки». Логика — см. references/fill_architecture.md.

Возможности v2:
  - заполнение шапки и переменных столбцов (План/Факт/Ответственный/Сроки/Примечание);
  - переопределение названия конфиг-строк (4.3.1–4.3.3, 6.1);
  - вставка дополнительных строк (1.1, 6.2, 6.3 …) с сохранением форматирования;
  - режимы: 'plan' (Факт пустой) и 'report' (Факт заполняется).

Использование: заполни DATA → `python3 fill_karta.py`
Зависимость: pip install python-docx --break-system-packages
"""
import os, re, copy
from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, '..', 'references', 'template_karta.docx')

# =====================================================================
# DATA — ЗАПОЛНИ ПОД ПРОЕКТ (схема — в references/fill_architecture.md)
# =====================================================================
DATA = {
    'gp':       'ТОО «........»',
    'irn':      'DPxxxxxxxx',
    'irn_name': 'DPxxxxxxxx «Наименование проекта»',
    'dogovor':  '№___ от «__» _______ 202_ года',
    'mode':     'plan',          # 'plan' | 'report'

    'rows': {
        '1':     {'plan': '', 'fact': '', 'otv': 'Руководитель проекта', 'srok': ''},
        '2':     {'plan': 'Не запланировано', 'otv': '-', 'srok': '-'},
        '3':     {'plan': 'Исполнить', 'otv': 'Технолог', 'srok': ''},
        '4.1':   {'plan': '', 'otv': 'Специалист по коммерциализации', 'srok': ''},
        '4.2':   {'plan': '', 'otv': 'Специалист по коммерциализации', 'srok': ''},
        '4.3.1': {'name': '', 'plan': 'Исполнить', 'otv': 'Специалист по коммерциализации', 'srok': ''},
        '4.3.2': {'name': '', 'plan': 'Исполнить', 'otv': 'Специалист по коммерциализации', 'srok': ''},
        '4.3.3': {'name': '', 'plan': 'Исполнить', 'otv': 'Специалист по коммерциализации', 'srok': ''},
        '4.4':   {'plan': 'Не менее 1 договора', 'otv': 'Специалист по коммерциализации', 'srok': ''},
        '4.5':   {'plan': '', 'otv': 'Руководитель проекта', 'srok': ''},
        '4.6.1': {'plan': '', 'otv': 'Специалист по коммерциализации', 'srok': ''},
        '4.6.2': {'plan': '-', 'otv': '-', 'srok': '-'},
        '4.7':   {'plan': '', 'otv': 'Руководитель проекта', 'srok': ''},
        '5.1':   {'plan': '', 'otv': 'Руководитель проекта', 'srok': ''},
        '5.2':   {'plan': '', 'otv': 'Бухгалтер проекта', 'srok': ''},
        '5.3':   {'plan': '-', 'otv': '-', 'srok': '-'},
        '6.1':   {'name': '', 'plan': '', 'otv': '', 'srok': ''},
    },

    # Вставляемые строки (нет в каркасе). 'after' — № строки, после которой вставить.
    'extra_rows': [
        # {'after': '1',   'num': '1.1', 'name': '...', 'plan': '1', 'otv': 'Руководитель проекта', 'srok': '...'},
        # {'after': '6.1', 'num': '6.2', 'name': '...', 'plan': 'Исполнить', 'otv': '...', 'srok': '...'},
    ],
}

COLMAP = {3: 'plan', 4: 'fact', 5: 'otv', 6: 'srok', 7: 'prim'}


def set_cell(cell, text):
    """Записать текст в ячейку, сохранив форматирование первого run."""
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    p = cell.paragraphs[0]
    if not p.runs:
        p.add_run('')
    base = p.runs[0]
    for r in p.runs[1:]:
        r._element.getparent().remove(r._element)
    base.text = str(text)


def build_num_map(t):
    """№ → индекс строки (строки 0-1 — шапка таблицы, пропускаются)."""
    m = {}
    for ri, row in enumerate(t.rows):
        if ri < 2:
            continue
        num = row.cells[0].text.strip().rstrip('.')
        if num and num not in m:
            m[num] = ri
    return m


def insert_extra_rows(t, extra_rows):
    """Клонировать строку-донор ('after') и вставить следом новую строку."""
    for spec in extra_rows:
        num_map = build_num_map(t)
        anchor_ri = num_map.get(str(spec['after']).rstrip('.'))
        if anchor_ri is None:
            print('  ! строка-якорь №%s не найдена — вставка пропущена' % spec['after'])
            continue
        anchor_tr = t.rows[anchor_ri]._tr
        new_tr = copy.deepcopy(anchor_tr)
        anchor_tr.addnext(new_tr)
        # очистить переменные ячейки клонированной строки
        new_row = [r for r in t.rows if r._tr is new_tr][0]
        set_cell(new_row.cells[0], spec.get('num', ''))
        set_cell(new_row.cells[1], spec.get('name', ''))
        for ci in COLMAP:
            set_cell(new_row.cells[ci], '')


def set_header(doc, key, value):
    token = '{{%s}}' % key
    for p in doc.paragraphs:
        if token in p.text:
            full = p.text.replace(token, value)
            p.runs[0].text = full
            for r in p.runs[1:]:
                r._element.getparent().remove(r._element)
            return


def main():
    d = Document(TEMPLATE)
    set_header(d, 'GP', DATA['gp'])
    set_header(d, 'IRN_NAME', DATA['irn_name'])
    set_header(d, 'DOGOVOR', DATA['dogovor'])

    t = d.tables[0]
    insert_extra_rows(t, DATA.get('extra_rows', []))

    report = DATA.get('mode') == 'report'
    num_map = build_num_map(t)

    # каркасные + конфиг-строки
    for num, vals in DATA['rows'].items():
        ri = num_map.get(num)
        if ri is None:
            print('  ! строка № %s не найдена — пропущена' % num)
            continue
        if vals.get('name'):
            set_cell(t.rows[ri].cells[1], vals['name'])
        for ci, key in COLMAP.items():
            if key == 'fact' and not report:
                continue
            if key in vals and vals[key] != '':
                set_cell(t.rows[ri].cells[ci], vals[key])

    # переменные значения вставленных строк
    num_map = build_num_map(t)
    for spec in DATA.get('extra_rows', []):
        ri = num_map.get(str(spec.get('num', '')).rstrip('.'))
        if ri is None:
            continue
        for ci, key in COLMAP.items():
            if key == 'fact' and not report:
                continue
            if key in spec and spec[key] != '':
                set_cell(t.rows[ri].cells[ci], spec[key])

    safe = re.sub(r'[^A-Za-z0-9_-]', '', DATA['irn']) or 'project'
    out = os.path.join(os.getcwd(), 'Карта_внедрения_%s.docx' % safe)
    d.save(out)
    print('Сохранено:', out)


if __name__ == '__main__':
    main()
