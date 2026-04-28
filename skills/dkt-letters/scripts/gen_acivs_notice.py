"""
gen_acivs_notice.py — Генератор уведомления грантополучателю по АЦИВС
АО «Фонд науки» / ДКТ

Структура шаблона (индексы параграфов):
  P[0]  — логотип Фонда (изображение, НЕ трогать)
  P[1]  — адресная строка (НЕ трогать)
  P[4]  — получатель строка 1 (жирный)
  P[5]  — получатель строка 2 (жирный)
  P[8]  — вводный абзац
  P[9]  — пункт о праве на финансирование следующего этапа
  P[10] — пункт об удержании экономии из финансирования следующего этапа
  P[11] — «Так, согласно заключению АЦИВС по Проекту:»
  P[12] — bullet 1: мероприятия КП
  P[13] — bullet 2: экономия
  P[14] — bullet 3: нецелевое
  P[15] — bullet 4: со-финансирование
  P[16] — «На основании вышеизложенного»
  P[17] — пункт 1) подписанное заключение АЦИВС
  P[19] — «Приложение:»
  P[23] — подписант (Управляющий директор / ФИО)
  P[30] — исполнитель строка 1
  P[31] — исполнитель строка 2

Зависимость: pip install python-docx --break-system-packages
"""

import copy, glob
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt
from docx.text.paragraph import Paragraph

# ===================================================================
#  ПАРАМЕТРЫ ПРОЕКТА — заполнить из договора ГП
# ===================================================================

IRN           = "DP23692342"
STAGE         = 2
STAGE_RU      = "2-й"
PERIOD        = "январь–декабрь 2025 г."
NEXT_STAGE_OF = "3-го"

# Получатель (правый блок, жирный)
GP_LINE1 = "Сельскохозяйственный производственный кооператив"
GP_LINE2 = "«MELON juice Co»"   # "" если однострочно

# Данные договора
CONTRACT_NUM  = "107"
CONTRACT_DATE = "«24» сентября 2024 года"
PROJECT_NAME  = (
    "«Коммерциализация инновационных технологий производства "
    "функциональных овощных соков с высоким содержанием полезных компонентов»"
)

# Пункты договора (взять из конкретного Договора!)
CLAUSE_NEXT_STAGE = "3.4"   # право на финансирование следующего этапа
CLAUSE_ECONOMY    = "3.7"   # удержание экономии из финансирования следующего этапа

# Подписант
SIGNATORY_TITLE = "Управляющий директор"
SIGNATORY_NAME  = "А. Рахимжанов"

# Исполнитель
EXECUTOR_NAME   = "Исп.: А. Снадин"
EXECUTOR_PHONE  = "8 (7172) 57-50-08 / 57-50-15"

# Экономия гранта (None — если нет)
ECO_SUM = (
    "3 158 414 (три миллиона сто пятьдесят восемь тысяч четыреста "
    "четырнадцать) тенге 00 тиын"
)

# Нецелевое использование
NECELEV_NOT_FOUND = True   # True -> «не выявлено»

# Со-финансирование
SOFIN_SHORTFALL = True
SOFIN_PLAN_SUM  = "88 000 000,00 (восемьдесят восемь миллионов)"
SOFIN_NEDOSTACH = (
    "39 490 300,00 (тридцать девять миллионов четыреста девяносто "
    "тысяч триста)"
)
SOFIN_ACCOUNT = "KZ558562203140010240"
SOFIN_BANK    = "АО «БанкЦентрКредит»"

# Пункты договора для блока со-финансирования
CLAUSE_SOFIN_DEPOSIT  = "2.16"   # обеспечить вложение средств на р/с
CLAUSE_SOFIN_SPENDING = "2.8"    # расходовать исключительно на цели проекта

# Счета ГП для запроса выписки
ACCOUNTS = [
    {"type": "эскроу-счёт",
     "num": "KZ168562203140864348",
     "bank": "АО «БанкЦентрКредит»"},
    {"type": "текущий счёт",
     "num": "KZ668562203141697551",
     "bank": "АО «БанкЦентрКредит»"},
    {"type": "счёт со-финансирования",
     "num": SOFIN_ACCOUNT,
     "bank": SOFIN_BANK},
]

# АЦИВС
ACIVS_PAGES = 13

# Пути
TEMPLATE = ""
OUTPUT   = f"Уведомление_АЦИВС{STAGE}этап_{IRN}.docx"

if not TEMPLATE:
    candidates = (
        glob.glob("/sessions/*/mnt/uploads/Уведомление_АЦИВС*этап_*.docx") +
        glob.glob("/sessions/*/mnt/**/*Уведомление_АЦИВС*этап*.docx", recursive=True)
    )
    if not candidates:
        raise FileNotFoundError("Шаблон не найден. Загрузите Уведомление_АЦИВС*этап_*.docx")
    TEMPLATE = candidates[0]
    print(f"Шаблон: {TEMPLATE}")

# ===================================================================
#  ХЕЛПЕРЫ
# ===================================================================

def clear_runs(para):
    p = para._p
    for tag in ("w:r", "w:ins", "w:hyperlink", "w:del"):
        for el in p.findall(qn(tag)):
            p.remove(el)

def add_run(para, text, bold=False, italic=False, sz=12):
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(sz)
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rFonts.set(qn(attr), "Times New Roman")
    return run

def clone_ppr(para):
    pPr = para._p.find(qn("w:pPr"))
    return copy.deepcopy(pPr) if pPr is not None else None

def insert_after(ref_para, text_parts, pPr_clone=None):
    new_p = OxmlElement("w:p")
    if pPr_clone is not None:
        new_p.append(copy.deepcopy(pPr_clone))
    ref_para._p.addnext(new_p)
    para = Paragraph(new_p, ref_para._p.getparent())
    for part in text_parts:
        add_run(para, part["text"],
                bold=part.get("bold", False),
                italic=part.get("italic", False),
                sz=part.get("sz", 12))
    return para

# ===================================================================
#  БЛОКИ ТЕКСТА
# ===================================================================

PARA1 = (
    f"АО «Фонд науки» (далее – Фонд) направляет для ознакомления и подписания заключение "
    f"анализа целевого использования выделенных средств (далее – АЦИВС) по промежуточному "
    f"отчёту за {STAGE_RU} этап проекта ИРН {IRN} "
    f"{PROJECT_NAME} (далее – Проект), реализуемого в "
    f"соответствии с Договором о предоставлении гранта на коммерциализацию результатов "
    f"научной и (или) научно-технической деятельности № {CONTRACT_NUM} от {CONTRACT_DATE} "
    f"(далее – Договор)."
)
PARA9 = (
    f"В соответствии с пунктом {CLAUSE_NEXT_STAGE} Договора, Грантополучатель получает "
    f"право на получение финансирования следующего этапа на основании положительного "
    f"решения ННС об утверждении промежуточного отчёта."
)
PARA10 = (
    f"Также, согласно пункту {CLAUSE_ECONOMY} Договора, в случае экономии денежных средств "
    f"по предыдущему этапу, сумма финансирования последующего этапа будет уменьшена на "
    f"соответствующую денежную сумму экономии по предыдущему этапу."
)
PRILOZHENIE = (
    f"Приложение: Заключение анализа целевого использования выделенных средств (АЦИВС) "
    f"по промежуточному отчёту за {STAGE_RU} этап по проекту {IRN} "
    f"на {ACIVS_PAGES} листах."
)
SIGNATORY_LINE = SIGNATORY_TITLE + " " * max(1, 60 - len(SIGNATORY_TITLE)) + SIGNATORY_NAME

if ECO_SUM:
    eco_n = "экономия средств гранта составила "
    eco_b = ECO_SUM
    eco_e = f". Сумма экономии будет удержана из объёма финансирования {NEXT_STAGE_OF} этапа Проекта;"
else:
    eco_n = "экономия средств гранта "
    eco_b = "отсутствует"
    eco_e = ";"

net_end = ";" if SOFIN_SHORTFALL else "."
net_n   = "нецелевое и незапланированное использование средств гранта " if NECELEV_NOT_FOUND else "нецелевое использование средств гранта "
net_b   = "не выявлено" if NECELEV_NOT_FOUND else "выявлено"

# ===================================================================
#  ГЕНЕРАЦИЯ
# ===================================================================

doc   = Document(TEMPLATE)
paras = doc.paragraphs

# P[0] — логотип (НЕ трогаем); P[1] — адрес (НЕ трогаем)

clear_runs(paras[4]); add_run(paras[4], GP_LINE1, bold=True, sz=12)
if GP_LINE2:
    clear_runs(paras[5]); add_run(paras[5], GP_LINE2, bold=True, sz=12)

clear_runs(paras[8]);  add_run(paras[8],  PARA1,  sz=12)
clear_runs(paras[9]);  add_run(paras[9],  PARA9,  sz=12)
clear_runs(paras[10]); add_run(paras[10], PARA10, sz=12)

clear_runs(paras[12])
add_run(paras[12], f"все мероприятия Календарного плана {STAGE_RU} этапа ({PERIOD}) ", sz=12)
add_run(paras[12], "выполнены в полном объёме;", bold=True, sz=12)

clear_runs(paras[13])
add_run(paras[13], eco_n, sz=12); add_run(paras[13], eco_b, bold=True, sz=12); add_run(paras[13], eco_e, sz=12)

clear_runs(paras[14])
add_run(paras[14], net_n, sz=12); add_run(paras[14], net_b, bold=True, sz=12); add_run(paras[14], net_end, sz=12)

pPr_body = clone_ppr(paras[8])
clear_runs(paras[15])

if not SOFIN_SHORTFALL:
    add_run(paras[15], "обязательства по со-финансированию исполнены ", sz=12)
    add_run(paras[15], "в полном объёме.", bold=True, sz=12)
else:
    add_run(paras[15],
        f"В соответствии с пунктом {CLAUSE_SOFIN_DEPOSIT} Договора, Грантополучатель обязан:\n"
        f"обеспечить вложение денежных средств со-финансирования на расчетный счет "
        f"Грантополучателя № {SOFIN_ACCOUNT} в {SOFIN_BANK} "
        f"в объеме, предусмотренным настоящим Договором.", sz=12)
    p_k = insert_after(paras[15],
        [{"text":
            f"Кроме того, в соответствии с пунктом {CLAUSE_SOFIN_SPENDING} Договора, "
            f"средства со-финансирования расходуются в полном объеме и исключительно "
            f"на цели и задачи Проекта на каждом этапе его реализации.", "sz": 12}],
        pPr_clone=pPr_body)
    insert_after(p_k,
        [{"text":
            f"Далее, согласно заключению АЦИВС по Проекту, из суммы со-финансирования, "
            f"предусмотренной на {STAGE_RU} этапе, в размере {SOFIN_PLAN_SUM} тенге, "
            f"выявлено неисполнение обязательств на сумму {SOFIN_NEDOSTACH} тенге.", "sz": 12}],
        pPr_clone=pPr_body)

all_paras = doc.paragraphs
p16_idx = next(i for i, p in enumerate(all_paras) if p.text.startswith("На основании вышеизложенного"))
p17 = all_paras[p16_idx + 1]

if SOFIN_SHORTFALL:
    p2 = insert_after(p17,
        [{"text": f"2)     документы, подтверждающие исполнение обязательств по со-финансированию на сумму {SOFIN_NEDOSTACH} тенге.", "sz": 12}],
        pPr_clone=clone_ppr(p17))
    accs = "; ".join(f"{a['type']} № {a['num']} в {a['bank']}" for a in ACCOUNTS)
    insert_after(p2,
        [{"text": f"3)     выписки о движении средств с отметкой банка за отчетный период по следующим счетам Грантополучателя: {accs}.", "sz": 12}],
        pPr_clone=clone_ppr(p17))

all_paras = doc.paragraphs
p_pril = next(p for p in all_paras if p.text.startswith("Приложение:"))
clear_runs(p_pril); add_run(p_pril, PRILOZHENIE, italic=True, sz=10)

all_paras = doc.paragraphs
p_sign = next((p for p in all_paras if "Управляющий директор" in p.text or "Председатель Правления" in p.text), None)
if p_sign:
    clear_runs(p_sign); add_run(p_sign, SIGNATORY_LINE, sz=12)

all_paras = doc.paragraphs
p_exec1 = next((p for p in all_paras if p.text.startswith("Исп.:")), None)
if p_exec1:
    clear_runs(p_exec1); add_run(p_exec1, EXECUTOR_NAME, sz=10)
    idx = all_paras.index(p_exec1)
    if idx + 1 < len(all_paras):
        clear_runs(all_paras[idx+1]); add_run(all_paras[idx+1], EXECUTOR_PHONE, sz=10)

doc.save(OUTPUT)
print(f"Сохранено: {OUTPUT}")
